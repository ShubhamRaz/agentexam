import os
import re
import fitz  # PyMuPDF
import docx
from typing import Optional
import logging

logger = logging.getLogger(__name__)

class OCREngine:
    def extract_text_from_image(self, file_path: str) -> str:
        """
        Extract text from an image file (PNG/JPG/JPEG) or a scanned PDF using pytesseract + Pillow.
        Falls back to empty string if either library is not installed or OCR fails.
        """
        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            logger.warning(
                "pytesseract or Pillow not installed — OCR unavailable. "
                "Install with: pip install pytesseract pillow"
            )
            return ""

        text = ""
        ext = file_path.split(".")[-1].lower() if "." in file_path else ""

        try:
            if ext in ["png", "jpg", "jpeg"]:
                img = Image.open(file_path)
                page_text = pytesseract.image_to_string(img)
                if page_text:
                    text += page_text + "\n"
            else:
                # PDF fallback via PyMuPDF rendering
                with fitz.open(file_path) as doc:
                    for page in doc:
                        # Render page to image at 300 dpi
                        mat = fitz.Matrix(300 / 72, 300 / 72)
                        pix = page.get_pixmap(matrix=mat)
                        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                        page_text = pytesseract.image_to_string(img)
                        if page_text:
                            text += page_text + "\n"
        except Exception as e:
            logger.error(f"OCR failed for {file_path}: {e}")

        return text


class DocumentProcessor:
    def __init__(self, ocr_engine: Optional[OCREngine] = None):
        self.ocr_engine = ocr_engine or OCREngine()
        
    def extract_text(self, file_path: str, file_type: str) -> str:
        text = ""
        file_type_lower = file_type.lower()
        if file_type_lower == "pdf":
            text = self._extract_pdf(file_path)
        elif file_type_lower in ["docx", "doc"]:
            text = self._extract_docx(file_path)
        elif file_type_lower == "txt":
            text = self._extract_txt(file_path)
        elif file_type_lower in ["png", "jpg", "jpeg"]:
            text = self._extract_image(file_path)
        else:
            raise ValueError(f"Unsupported file type for extraction: {file_type}")
            
        return self._clean_text(text)
        
    def _extract_pdf(self, file_path: str) -> str:
        text = ""
        try:
            with fitz.open(file_path) as doc:
                for page in doc:
                    page_text = page.get_text()
                    if page_text:
                        text += page_text + "\n"
                        
            # If text is extremely short or empty, it might be a scanned PDF.
            if len(text.strip()) < 50:
                # Fallback to OCR
                text = self.ocr_engine.extract_text_from_image(file_path)
                
        except Exception as e:
            raise RuntimeError(f"PDF extraction failed: {str(e)}")
            
        return text

    def _extract_docx(self, file_path: str) -> str:
        text = ""
        try:
            doc = docx.Document(file_path)
            # Extract paragraphs
            for p in doc.paragraphs:
                if p.text.strip():
                    text += p.text + "\n"
                    
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        text += " | ".join(row_data) + "\n"
        except Exception as e:
            raise RuntimeError(f"DOCX extraction failed: {str(e)}")
            
        return text

    def _extract_txt(self, file_path: str) -> str:
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
        except Exception as e:
            raise RuntimeError(f"TXT extraction failed: {str(e)}")

    def _extract_image(self, file_path: str) -> str:
        return self.ocr_engine.extract_text_from_image(file_path)
        
    def _clean_text(self, text: str) -> str:
        # Normalize newlines
        text = re.sub(r'\r\n', '\n', text)
        # Remove multiple consecutive blank lines
        text = re.sub(r'\n{3,}', '\n\n', text)
        # Replace multiple spaces with a single space
        text = re.sub(r'[ \t]{2,}', ' ', text)
        return text.strip()
